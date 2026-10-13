import streamlit as st
import os
import io
import tempfile
import zipfile
import glob
import pandas as pd
from datetime import datetime
from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from PIL import Image
import openpyxl
import warnings

# Suppress openpyxl warnings
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

st.set_page_config(page_title="Specification Appraisal Generator", layout="centered")
st.title("📋 Specification Appraisal Generator")
st.write("Upload your files below to group POs by description and generate pre-populated Word appraisals.")

# File Uploaders
csv_file = st.file_uploader("1. Upload CSV Order List", type=["csv"])
template_file = st.file_uploader("2. Upload Word Template (.docx)", type=["docx"])
excel_files = st.file_uploader("3. Upload Excel Sheet(s) with Locked Images (.xlsx)", type=["xlsx"], accept_multiple_files=True)

def find_paragraphs_in_shapes(doc):
    paragraphs = []
    try:
        raw_element = doc.element
        for txbx in raw_element.iterfind('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent'):
            for p_elem in txbx.iterfind('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                paragraphs.append(Paragraph(p_elem, doc))
        for shape in raw_element.iterfind('.//{urn:schemas-microsoft-com:vml}textbox'):
            for p_elem in shape.iterfind('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                p_obj = Paragraph(p_elem, doc)
                if p_obj not in paragraphs:
                    paragraphs.append(p_obj)
    except Exception as e:
        st.write(f"Note scanning deep shapes: {e}")
    return paragraphs

def map_images_from_locked_cells(target_pos, excel_dir):
    po_image_map = {}
    excel_paths = glob.glob(os.path.join(excel_dir, "*.xlsx"))
    excel_paths = [f for f in excel_paths if not os.path.basename(f).startswith("~$")]

    for file_path in excel_paths:
        try:
            df_excel = pd.read_excel(file_path, header=None)
            row_to_po_map = {}
            for idx, row in df_excel.iterrows():
                excel_row_num = idx + 1
                row_tokens = [str(val).strip().split('.0')[0] for val in row.values if pd.notna(val)]
                for token in row_tokens:
                    if token in target_pos:
                        row_to_po_map[excel_row_num] = token
                        break

            if not row_to_po_map:
                continue

            wb = openpyxl.load_workbook(file_path, data_only=True, keep_links=False)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                if not hasattr(ws, '_images') or not ws._images:
                    continue
                
                for img_obj in ws._images:
                    row_idx = None
                    if img_obj.anchor:
                        if hasattr(img_obj.anchor, 'row') and img_obj.anchor.row is not None:
                            row_idx = img_obj.anchor.row
                        elif hasattr(img_obj.anchor, '_from') and hasattr(img_obj.anchor._from, 'row'):
                            row_idx = img_obj.anchor._from.row
                    
                    if row_idx is not None:
                        img_row = row_idx + 1
                        matched_po = row_to_po_map.get(img_row)
                        if matched_po:
                            try:
                                po_image_map[matched_po] = Image.open(img_obj.ref).copy()
                            except Exception:
                                pass
            wb.close()
        except Exception:
            pass
            
    return po_image_map

if st.button("⚡ Generate Appraisals", type="primary"):
    if not csv_file or not template_file:
        st.error("Please upload both the CSV Order List and the Word Template.")
    else:
        with st.spinner("Processing documents and mapping images..."):
            with tempfile.TemporaryDirectory() as temp_dir:
                # Save input files to temporary directory
                csv_path = os.path.join(temp_dir, "data.csv")
                template_path = os.path.join(temp_dir, "template.docx")
                excel_dir = os.path.join(temp_dir, "excel_files")
                os.makedirs(excel_dir, exist_ok=True)

                with open(csv_path, "wb") as f:
                    f.write(csv_file.getbuffer())
                with open(template_path, "wb") as f:
                    f.write(template_file.getbuffer())

                if excel_files:
                    for ef in excel_files:
                        with open(os.path.join(excel_dir, ef.name), "wb") as f:
                            f.write(ef.getbuffer())

                df = pd.read_csv(csv_path)
                today_date = datetime.today().strftime('%d/%m/%Y')
                
                target_pos = set()
                for _, row in df.iterrows():
                    po = str(row.get('PurchaseOrderNumber', '')).strip().split('.0')[0]
                    if po and po != 'nan':
                        target_pos.add(po)

                excel_po_images = map_images_from_locked_cells(target_pos, excel_dir)

                qa_mapping = {
                    'ORDER NUMBER': 'PurchaseOrderNumber',
                    'ITEM NUMBER': 'DeptItemCode',
                    'Supplier Style Number': 'SupplierStyleNo',
                    'SELLING PRICE': 'SellingPrice',
                    'DESCRIPTION': 'ItemDesc',
                    'UNITS': 'PurchQty',
                    'PRIMARY COLOUR': 'SupplierColor',
                    'SIZE(mm)': 'Size'
                }

                zip_buffer = io.BytesIO()
                total_generated = 0

                with zipfile.ZipFile(zip_buffer, "w") as zip_file:
                    grouped = df.groupby('ItemDesc', dropna=False)

                    for item_desc, group in grouped:
                        item_desc_str = str(item_desc) if pd.notna(item_desc) else ""
                        first_row = group.iloc[0]
                        
                        po_list = []
                        for _, r in group.iterrows():
                            p_num = str(r.get('PurchaseOrderNumber', '')).split('.0')[0]
                            if p_num and p_num not in po_list:
                                po_list.append(p_num)

                        primary_po = po_list[0] if po_list else "UNKNOWN"
                        item_num = str(first_row.get('DeptItemCode', '')).split('.0')[0]
                        delivery_date = str(first_row.get('DeliveryDate', '')) if pd.notna(first_row.get('DeliveryDate', '')) else ""

                        output_filename = f"Specification_Appraisal_PO_{'_'.join(po_list)}.docx" if len(group) > 1 else f"Specification_Appraisal_PO_{primary_po}.docx"
                        
                        doc = Document(template_path)

                        # Header Table Population
                        if len(doc.tables) > 0:
                            header_table = doc.tables[0]
                            for row in header_table.rows:
                                for c_idx in range(len(row.cells) - 1):
                                    label_text = row.cells[c_idx].text.strip()
                                    blank_cell = row.cells[c_idx + 1]
                                    
                                    if label_text == "DATE OF SUBMISION":
                                        blank_cell.text = today_date
                                    elif label_text == "DELIVERY DATE":
                                        blank_cell.text = delivery_date
                                    elif label_text == "DESCRIPTION":
                                        blank_cell.text = item_desc_str
                                    elif label_text == "ORDER NUMBER":
                                        blank_cell.text = primary_po if len(group) == 1 else ", ".join(po_list)
                                    elif label_text == "ITEM NUMBER":
                                        blank_cell.text = item_num

                        # Image Insertion
                        excel_image = None
                        for po in po_list:
                            if po in excel_po_images:
                                excel_image = excel_po_images[po]
                                break

                        if excel_image:
                            image_placeholder = "ITEM PICTURE INCLUDING INSERTS"
                            temp_img_path = os.path.join(temp_dir, f"temp_{primary_po}.jpg")
                            
                            img = excel_image
                            if img.mode in ('RGBA', 'LA') or (img.mode == 'P' and 'transparency' in img.info):
                                img = img.convert('RGB')
                            img.save(temp_img_path, 'JPEG', quality=60, optimize=True)

                            inserted = False
                            for paragraph in find_paragraphs_in_shapes(doc):
                                if image_placeholder in paragraph.text:
                                    paragraph.text = ""
                                    paragraph.add_run().add_picture(temp_img_path, width=Inches(2.2))
                                    inserted = True
                                    break

                            if not inserted:
                                for paragraph in doc.paragraphs:
                                    if image_placeholder in paragraph.text:
                                        paragraph.text = ""
                                        paragraph.add_run().add_picture(temp_img_path, width=Inches(2.2))
                                        inserted = True
                                        break

                            if not inserted:
                                for table in doc.tables:
                                    for r in table.rows:
                                        for cell in r.cells:
                                            if image_placeholder in cell.text:
                                                p = cell.paragraphs[0]
                                                p.text = ""
                                                p.add_run().add_picture(temp_img_path, width=Inches(2.2))
                                                inserted = True
                                                break
                                        if inserted: break

                        # QA Table Row Population
                        qa_table = None
                        for table in doc.tables:
                            if len(table.rows) > 0:
                                table_headers = [c.text.strip() for c in table.rows[0].cells]
                                if 'ORDER NUMBER' in table_headers and 'ITEM NUMBER' in table_headers:
                                    qa_table = table
                                    break

                        if qa_table:
                            headers = [c.text.strip() for c in qa_table.rows[0].cells]
                            default_row_values = [cell.text.strip() for cell in qa_table.rows[1].cells] if len(qa_table.rows) > 1 else []

                            for g_idx, (_, csv_row) in enumerate(group.iterrows()):
                                target_row_idx = g_idx + 1
                                row_to_fill = qa_table.rows[target_row_idx] if target_row_idx < len(qa_table.rows) else qa_table.add_row()

                                for col_idx, header in enumerate(headers):
                                    if header in qa_mapping:
                                        csv_col_name = qa_mapping[header]
                                        if csv_col_name in df.columns:
                                            val = csv_row[csv_col_name]
                                            cell_val = str(val).split('.0')[0] if pd.notna(val) else ""
                                            row_to_fill.cells[col_idx].text = cell_val
                                    else:
                                        if header in df.columns and pd.notna(csv_row[header]):
                                            row_to_fill.cells[col_idx].text = str(csv_row[header])
                                        elif col_idx < len(default_row_values):
                                            row_to_fill.cells[col_idx].text = default_row_values[col_idx]

                        # Save output into Zip Archive
                        doc_io = io.BytesIO()
                        doc.save(doc_io)
                        zip_path = output_filename if excel_image else f"Missing_Images/{output_filename}"
                        zip_file.writestr(zip_path, doc_io.getvalue())
                        total_generated += 1

                st.success(f"Generated {total_generated} appraisals!")
                st.download_button(
                    label="📥 Download All Appraisals (.ZIP)",
                    data=zip_buffer.getvalue(),
                    file_name="Generated_Appraisals.zip",
                    mime="application/zip",
                    type="primary"
                )
